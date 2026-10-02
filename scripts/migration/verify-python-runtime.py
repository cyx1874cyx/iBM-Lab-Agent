"""P1 local scientific-runtime smoke; no licensed application or remote service call."""
import argparse, hashlib, importlib, importlib.metadata, json, re, sys, zipfile
from pathlib import Path

parser=argparse.ArgumentParser(); parser.add_argument('--output',required=True); args=parser.parse_args()
out=Path(args.output).resolve();out.mkdir(parents=True,exist_ok=True)
root=Path(__file__).resolve().parents[2]
assert sys.version_info[:3]==(3,12,11),sys.version
lock=root/'python/requirements.lock'; versions={}
for line in lock.read_text(encoding='utf-8').splitlines():
    match=re.fullmatch(r'([\w.-]+)==([^\s#]+)',line.strip())
    if match:
        name,expected=match.groups();actual=importlib.metadata.version(name);assert actual==expected,(name,actual,expected);versions[name]=actual
for name in ['numpy','scipy','pandas','matplotlib','nmrglue','fitz','pptx','openpyxl','mcp','origin_mcp','mnova_mcp']:
    importlib.import_module(name)
from markitdown import MarkItDown
from pptx import Presentation
from openpyxl import Workbook
import nmrglue as ng
import numpy as np
import fitz

text='P1 scientific runtime verification'
docx=out/'probe.docx'
with zipfile.ZipFile(docx,'w') as z:
    z.writestr('[Content_Types].xml','<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    z.writestr('_rels/.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    z.writestr('word/document.xml',f'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>')
presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[0]);slide.shapes.title.text=text; presentation.save(out/'probe.pptx')
book=Workbook();book.active.append([text,3.12]);book.save(out/'probe.xlsx')
converted=[]
for name in ['probe.docx','probe.pptx','probe.xlsx']:
    markdown=MarkItDown().convert(out/name).text_content;assert text in markdown,(name,markdown)
    (out/(name+'.md')).write_text(markdown,encoding='utf-8');converted.append(name)
document=fitz.open();page=document.new_page();page.insert_text((72,72),text);document.save(out/'probe.pdf');document.close()
with fitz.open(out/'probe.pdf') as document:assert text in document[0].get_text()
fid=np.exp(2j*np.pi*0.125*np.arange(1024))*np.exp(-np.arange(1024)/150)
spectrum=ng.proc_base.fft(fid);assert np.isfinite(spectrum).all();assert np.max(np.abs(spectrum))>10
np.save(out/'nmr-spectrum.npy',spectrum)
report={'ok':True,'python':'.'.join(map(str,sys.version_info[:3])),'requirementsSha256':hashlib.sha256(lock.read_bytes()).hexdigest(),'lockedPackages':versions,
        'checks':['all-locked-versions-match','scientific-and-mcp-imports','docx-pptx-xlsx-conversion','pdf-create-and-extract','nmrglue-numpy-fft'],
        'conversionInputs':converted,'scope':'Local foundations only; no Origin/Mnova application, institution browser or LLM invocation'}
(out/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps({'ok':True,'lockedPackages':len(versions),'checks':report['checks']}))
